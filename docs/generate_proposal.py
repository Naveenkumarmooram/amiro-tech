from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
PUBLIC = ROOT / 'public'
OUTPUT = ROOT / 'docs' / 'Amiro_LIMS_Budgetary_Commercial_Proposal_TUV_Rheinland_Professional.docx'
LOGO = ROOT / 'docs' / 'amiro-logo-proposal.png'

NAVY = '202A36'
GOLD = 'F4B400'
LIGHT = 'F4F6F8'
MID = '5B6570'
BORDER = 'D9DEE3'


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn('w:shd'))
    if element is None:
        element = OxmlElement('w:shd')
        properties.append(element)
    element.set(qn('w:fill'), fill)


def set_cell_border(cell, **kwargs):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        properties.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            tag = 'w:{}'.format(edge)
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ['val', 'sz', 'space', 'color']:
                if key in kwargs[edge]:
                    element.set(qn('w:{}'.format(key)), str(kwargs[edge][key]))


def set_cell_text(cell, text, bold=False, color=NAVY, size=9.5):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = 'Aptos'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = 'Aptos Display'
    run.font.size = Pt(17 if level == 1 else 12)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    return paragraph


def add_body(document, text, size=10, color=MID, after=6):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    run.font.name = 'Aptos'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style='List Bullet')
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = 'Aptos'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(MID)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Amiro Tech Solutions  |  Confidential  |  ')
    run.font.name = 'Aptos'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(field)


def create_print_logo():
    image = Image.new('RGB', (1040, 300), 'white')
    draw = ImageDraw.Draw(image)
    mark = [(80, 235), (165, 42), (255, 235), (210, 235), (165, 112), (125, 235)]
    draw.polygon(mark, fill='#F4B400')
    draw.polygon([(168, 42), (255, 235), (210, 235), (145, 92)], fill='#2A2C31')
    draw.polygon([(150, 145), (185, 215), (220, 215), (190, 145)], fill='#2A2C31')
    draw.polygon([(165, 104), (214, 195), (116, 195)], fill='white')
    draw.polygon([(165, 147), (148, 180), (182, 180)], fill='#F4B400')
    try:
        word_font = ImageFont.truetype('C:/Windows/Fonts/segoeuib.ttf', 82)
        sub_font = ImageFont.truetype('C:/Windows/Fonts/segoeui.ttf', 34)
    except OSError:
        word_font = ImageFont.load_default()
        sub_font = ImageFont.load_default()
    draw.text((330, 84), 'Amiro Tech', font=word_font, fill='#2A2C31')
    draw.text((455, 185), 'S O L U T I O N S', font=sub_font, fill='#777F86')
    image.save(LOGO)


create_print_logo()

document = Document()
section = document.sections[0]
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = document.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10)
styles['Normal'].font.color.rgb = RGBColor.from_string(MID)

header = section.header
header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
header_run = header.paragraphs[0].add_run()
header_run.add_picture(str(LOGO), width=Inches(2.55))
header.paragraphs[0].paragraph_format.space_after = Pt(3)

footer = section.footer
add_page_number(footer.paragraphs[0])

# Title block
paragraph = document.add_paragraph()
paragraph.paragraph_format.space_before = Pt(16)
paragraph.paragraph_format.space_after = Pt(4)
run = paragraph.add_run('BUDGETARY COMMERCIAL PROPOSAL')
run.bold = True
run.font.name = 'Aptos Display'
run.font.size = Pt(25)
run.font.color.rgb = RGBColor.from_string(NAVY)
accent = paragraph.add_run('  /  LIMS IMPLEMENTATION')
accent.bold = True
accent.font.name = 'Aptos Display'
accent.font.size = Pt(11)
accent.font.color.rgb = RGBColor.from_string(GOLD)

paragraph = document.add_paragraph()
paragraph.paragraph_format.space_after = Pt(14)
run = paragraph.add_run('Submitted to TÜV Rheinland')
run.bold = True
run.font.name = 'Aptos'
run.font.size = Pt(15)
run.font.color.rgb = RGBColor.from_string(MID)

meta = document.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
meta.columns[0].width = Inches(1.65)
meta.columns[1].width = Inches(4.8)
metadata = [
    ('Prepared by', 'Amiro Tech Solutions'),
    ('Proposal type', 'Budgetary commercial proposal'),
    ('Issue date', date(2026, 8, 27).strftime('%d %B %Y')),
    ('Commercial validity', '30 days from the issue date'),
]
for row, (label, value) in zip(meta.rows, metadata):
    set_cell_text(row.cells[0], label.upper(), bold=True, color=NAVY, size=8.5)
    set_cell_text(row.cells[1], value, color=MID, size=9.5)
    shade(row.cells[0], LIGHT)
    for cell in row.cells:
        set_cell_border(cell, top={'val': 'single', 'sz': 5, 'color': BORDER}, bottom={'val': 'single', 'sz': 5, 'color': BORDER}, left={'val': 'single', 'sz': 5, 'color': BORDER}, right={'val': 'single', 'sz': 5, 'color': BORDER})

callout = document.add_table(rows=1, cols=2)
callout.alignment = WD_TABLE_ALIGNMENT.LEFT
callout.autofit = False
callout.columns[0].width = Inches(4.55)
callout.columns[1].width = Inches(1.9)
set_cell_text(callout.cell(0, 0), 'INDICATIVE INVESTMENT\nLIMS implementation for TÜV Rheinland', bold=True, color='FFFFFF', size=10.5)
set_cell_text(callout.cell(0, 1), '₹17-22 lakh', bold=True, color='202A36', size=17)
shade(callout.cell(0, 0), NAVY)
shade(callout.cell(0, 1), 'F4B400')
for cell in callout.rows[0].cells:
    set_cell_border(cell, top={'val': 'single', 'sz': 8, 'color': NAVY}, bottom={'val': 'single', 'sz': 8, 'color': NAVY}, left={'val': 'single', 'sz': 8, 'color': NAVY}, right={'val': 'single', 'sz': 8, 'color': NAVY})

add_body(document, 'Prepared for commercial evaluation and planning. Final scope and fixed commercials will follow requirements validation.', size=9, color=MID, after=0)
document.add_page_break()

add_heading(document, '1. Executive Summary')
add_body(document, 'Amiro Tech Solutions is pleased to submit this budgetary commercial proposal to TÜV Rheinland for the design and implementation of a Laboratory Information Management System (LIMS). The proposed solution is intended to provide a controlled digital environment for laboratory operations, equipment records, calibration, compliance evidence, reporting, and audit readiness.')
add_body(document, 'This proposal has been prepared at the budgetary stage to support planning and commercial evaluation. It is based on the scope currently understood and will be refined through a structured discovery and requirements-validation phase.')

add_heading(document, '2. Proposed Solution')
add_body(document, 'Amiro will deliver a secure, role-based LIMS platform configured around TÜV Rheinland\'s operating procedures and reporting needs. The implementation will prioritize traceability, data integrity, usability, and maintainability, with clear checkpoints for review and acceptance.')
for item in [
    'Centralized management of samples, tests, jobs, results, and laboratory records.',
    'Equipment register with calibration schedules, certificates, status, and lifecycle history.',
    'Role-based access, approvals, audit trails, and controlled record changes.',
    'Operational dashboards, management reports, exports, and configurable notifications.',
    'Document management for procedures, certificates, evidence, and supporting attachments.',
    'Deployment, training, user acceptance support, and handover documentation.',
]:
    add_bullet(document, item)

add_heading(document, '3. Indicative Scope and Budget')
add_body(document, 'The following figures are indicative estimates in Indian Rupees. The overall implementation budget is expected to be within the range of ₹17,00,000 to ₹22,00,000, subject to confirmation of detailed requirements, integrations, data volume, environments, and acceptance criteria.')

table = document.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [0.48, 3.42, 1.25, 1.25]
for cell, width in zip(table.rows[0].cells, widths):
    cell.width = Inches(width)
headers = ['No.', 'Work package', 'Budget low', 'Budget high']
for cell, text in zip(table.rows[0].cells, headers):
    set_cell_text(cell, text, bold=True, color='FFFFFF', size=9)
    shade(cell, NAVY)
    set_cell_border(cell, top={'val': 'single', 'sz': 5, 'color': NAVY}, bottom={'val': 'single', 'sz': 5, 'color': NAVY}, left={'val': 'single', 'sz': 5, 'color': NAVY}, right={'val': 'single', 'sz': 5, 'color': NAVY})

work_packages = [
    ('01', 'Discovery, requirements validation, and project planning', '₹1,50,000', '₹2,00,000'),
    ('02', 'Solution architecture, UX design, and technical foundation', '₹2,00,000', '₹2,50,000'),
    ('03', 'Core LIMS modules and workflow implementation', '₹6,00,000', '₹8,00,000'),
    ('04', 'Compliance controls, audit trail, dashboards, and reporting', '₹2,00,000', '₹3,00,000'),
    ('05', 'Integration, data migration, configuration, and environment setup', '₹2,00,000', '₹3,00,000'),
    ('06', 'Testing, UAT, deployment, training, and handover', '₹3,50,000', '₹3,50,000'),
]
for index, row_data in enumerate(work_packages):
    row = table.add_row()
    for cell, width in zip(row.cells, widths):
        cell.width = Inches(width)
    for cell, text in zip(row.cells, row_data):
        set_cell_text(cell, text, color=MID, size=8.8)
        set_cell_border(cell, top={'val': 'single', 'sz': 5, 'color': BORDER}, bottom={'val': 'single', 'sz': 5, 'color': BORDER}, left={'val': 'single', 'sz': 5, 'color': BORDER}, right={'val': 'single', 'sz': 5, 'color': BORDER})
    if index % 2 == 0:
        for cell in row.cells:
            shade(cell, 'FAFBFC')

total = table.add_row()
for cell, width in zip(total.cells, widths):
    cell.width = Inches(width)
set_cell_text(total.cells[0], '', color=NAVY, size=9)
set_cell_text(total.cells[1], 'TOTAL INDICATIVE IMPLEMENTATION BUDGET', bold=True, color=NAVY, size=8.8)
set_cell_text(total.cells[2], '₹17,00,000', bold=True, color=NAVY, size=9)
set_cell_text(total.cells[3], '₹22,00,000', bold=True, color=NAVY, size=9)
for cell in total.cells:
    shade(cell, 'FFF5D6')
    set_cell_border(cell, top={'val': 'single', 'sz': 8, 'color': GOLD}, bottom={'val': 'single', 'sz': 8, 'color': GOLD}, left={'val': 'single', 'sz': 5, 'color': BORDER}, right={'val': 'single', 'sz': 5, 'color': BORDER})

add_heading(document, '4. Delivery Approach and Indicative Timeline')
add_body(document, 'The engagement will be delivered through governed phases with regular reviews and documented approvals. A detailed schedule will be issued after discovery and confirmation of TÜV Rheinland\'s nominated stakeholders and dependencies.')
phase_table = document.add_table(rows=1, cols=3)
phase_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, text in zip(phase_table.rows[0].cells, ['Phase', 'Primary activities', 'Indicative duration']):
    set_cell_text(cell, text, bold=True, color='FFFFFF', size=9)
    shade(cell, NAVY)
phases = [
    ('1. Discover', 'Requirements, workflows, roles, data, integrations, and acceptance criteria', '2-3 weeks'),
    ('2. Design', 'Architecture, UX, security model, backlog, and delivery plan', '2-3 weeks'),
    ('3. Build', 'Configuration, development, integrations, and iterative demonstrations', '10-14 weeks'),
    ('4. Assure', 'System testing, security checks, UAT support, fixes, and readiness review', '3-4 weeks'),
    ('5. Launch', 'Production deployment, training, documentation, and handover', '1-2 weeks'),
]
for phase in phases:
    row = phase_table.add_row()
    for cell, text in zip(row.cells, phase):
        set_cell_text(cell, text, color=MID, size=8.8)
        set_cell_border(cell, top={'val': 'single', 'sz': 5, 'color': BORDER}, bottom={'val': 'single', 'sz': 5, 'color': BORDER}, left={'val': 'single', 'sz': 5, 'color': BORDER}, right={'val': 'single', 'sz': 5, 'color': BORDER})

add_heading(document, '5. Commercial Assumptions')
for item in [
    'This is a budgetary estimate, not a final fixed-price statement of work.',
    'Final scope, effort, milestones, payment schedule, and acceptance criteria will be documented after discovery.',
    'The estimate assumes timely access to relevant stakeholders, current process information, sample data, and required system documentation.',
    'Taxes, statutory levies, hosting, cloud services, software licenses, messaging charges, and third-party subscription fees are excluded unless specifically stated in the final agreement.',
    'Data migration effort assumes structured and reasonably clean source data. Extensive cleansing, digitization, or historical reconstruction may require separate estimation.',
    'Any material scope change, additional integration, new user group, or regulatory requirement introduced after approval will be assessed through change control.',
]:
    add_bullet(document, item)

add_heading(document, '6. Payment Milestones')
add_body(document, 'The following milestone structure is proposed for commercial planning and may be finalized during contracting:')
for item in [
    '20% on approval of the proposal and project commencement.',
    '20% on completion and approval of discovery and solution design.',
    '30% on completion of the core implementation milestone.',
    '20% on UAT readiness and deployment approval.',
    '10% on production handover and completion of agreed training and documentation.',
]:
    add_bullet(document, item)

add_heading(document, '7. Closing Statement')
add_body(document, 'Amiro Tech Solutions appreciates the opportunity to present this proposal to TÜV Rheinland. We would welcome a structured discussion with the relevant business and technical stakeholders to validate the operating model, prioritize the initial release, and convert this budgetary estimate into a detailed implementation proposal.')

closing = document.add_table(rows=2, cols=2)
closing.alignment = WD_TABLE_ALIGNMENT.LEFT
set_cell_text(closing.cell(0, 0), 'For Amiro Tech Solutions', bold=True, color=NAVY, size=9.5)
set_cell_text(closing.cell(0, 1), 'For TÜV Rheinland', bold=True, color=NAVY, size=9.5)
set_cell_text(closing.cell(1, 0), 'Authorized Signatory\nName: ____________________\nDate: ____________________', color=MID, size=9)
set_cell_text(closing.cell(1, 1), 'Acknowledged by\nName: ____________________\nDate: ____________________', color=MID, size=9)
for row in closing.rows:
    for cell in row.cells:
        set_cell_border(cell, top={'val': 'single', 'sz': 5, 'color': BORDER}, bottom={'val': 'single', 'sz': 5, 'color': BORDER}, left={'val': 'single', 'sz': 5, 'color': BORDER}, right={'val': 'single', 'sz': 5, 'color': BORDER})

# Keep the proposal compact and polished.
for paragraph in document.paragraphs:
    paragraph.paragraph_format.widow_control = True

document.core_properties.title = 'Amiro LIMS Budgetary Commercial Proposal - TÜV Rheinland'
document.core_properties.subject = 'Budgetary commercial proposal for LIMS implementation'
document.core_properties.author = 'Amiro Tech Solutions'
document.save(str(OUTPUT))
print(OUTPUT)
